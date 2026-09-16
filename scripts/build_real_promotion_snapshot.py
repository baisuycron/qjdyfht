"""Build the local read-only promotion snapshot used by the prototype.

The source workbook is a database export created on 2026-09-10.  Keeping the
conversion as a script makes the data provenance explicit and repeatable while
avoiding any database credentials in browser code.
"""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "marketing_activity_export_20260910" / "营销活动真实数据导出_接口中文表头_含玩法分类_剔除X元Y件_20260910.xlsx"
OUTPUT = ROOT / "public" / "data" / "real-promotion-activities.json"


def iso(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    next(iterator, None)
    return [row for row in iterator if row and isinstance(row[0], (int, float))]


workbook = load_workbook(SOURCE, read_only=True, data_only=True)
activity_rows = rows(workbook["活动主表"])
condition_rows = rows(workbook["条件明细"])
item_rows = rows(workbook["商品明细"])

conditions_by_activity: dict[int, list[dict]] = {}
for row in condition_rows:
    conditions_by_activity.setdefault(int(row[1]), []).append(
        {
            "id": int(row[0]),
            "rowno": int(row[3]),
            "wareid": row[4],
            "sumamt": row[5] or 0,
            "sumqty": row[6] or 0,
            "profitrate": row[7] or 0,
            "resprice": row[8],
            "createTime": iso(row[9]),
            "updateTime": iso(row[10]),
        }
    )

items_by_activity: dict[int, list[dict]] = {}
for row in item_rows:
    items_by_activity.setdefault(int(row[1]), []).append(
        {
            "id": int(row[0]),
            "rowno": int(row[3]),
            "pstid": row[4],
            "pstqty": row[5] or 0,
            "pstprice": row[6] or 0,
            "priceDisc": row[7] or 0,
            "createTime": iso(row[8]),
            "updateTime": iso(row[9]),
        }
    )

activities = []
for row in activity_rows:
    activity_id = int(row[0])
    raw = json.loads(row[17]) if row[17] else {}
    activity = {
        **raw,
        "id": activity_id,
        "pstplanno": str(row[1]),
        "marketingType": row[2],
        "promName": row[3],
        "playType": row[4],
        "busnos": row[5],
        "days": row[6],
        "weekdays": str(row[7] or "1111111"),
        "givenum": row[8] or 0,
        "giveprice": row[9] or 0,
        "givetype": row[10],
        "plannum": row[11],
        "repeatflag": row[12],
        "sumamt": row[13] or 0,
        "sumqty": row[14] or 0,
        "starttime": iso(row[15]),
        "endtime": iso(row[16]),
        "dbStatus": row[18],
        "createTime": iso(row[19]),
        "updateTime": iso(row[20]),
        "conditionItemList": conditions_by_activity.get(activity_id, []),
        "giftItemList": items_by_activity.get(activity_id, []),
    }
    activities.append(activity)

play_counts: dict[str, int] = {}
for activity in activities:
    play_counts[activity["playType"]] = play_counts.get(activity["playType"], 0) + 1

payload = {
    "meta": {
        "source": "MySQL数据库导出快照",
        "snapshotDate": "2026-09-10",
        "activityCount": len(activities),
        "conditionCount": len(condition_rows),
        "itemCount": len(item_rows),
        "excludedPlayTypes": ["X元Y件（任选）"],
        "playCounts": play_counts,
    },
    "activities": activities,
}

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"wrote {len(activities)} activities to {OUTPUT}")
